import io
from datetime import UTC, datetime

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.core.config import settings
from app.core.jobs import job_worker
from app.services.resume_processing import RESUME_PROCESS_JOB
from tests.conftest import (
    RESUME_SERVICE_LLM,
    _with_mongo,
    insert_job_description,
    make_pdf,
)

JD_SERVICE_LLM = "app.services.job_description_service.llm_service.generate_structured"
RESUME_PARSE_PAGES = "app.services.resume_processing.parse_pdf_pages"


async def fake_resume_extraction(prompt: str, schema, **kwargs):
    return schema(
        skills=["Python", "FastAPI"],
        experience=[],
        projects=[],
        education=[],
    )


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Jane Doe - Backend Engineer")
    document.add_paragraph("Built reliable Python services.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Company"
    table.cell(0, 1).text = "Acme"
    table.cell(1, 0).text = "Skills"
    table.cell(1, 1).text = "Python, FastAPI"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_list_resumes_requires_auth(client: TestClient) -> None:
    response = client.get("/api/resumes")
    assert response.status_code == 401


def test_upload_rejects_unsupported_extension(client: TestClient, auth_headers: dict) -> None:
    response = client.post(
        "/api/resumes/upload",
        headers=auth_headers,
        files={"file": ("resume.txt", b"hello world", "text/plain")},
    )
    assert response.status_code == 400
    assert "PDF and DOCX" in response.json()["detail"]


def test_upload_rejects_content_that_does_not_match_extension(
    client: TestClient, auth_headers: dict
) -> None:
    response = client.post(
        "/api/resumes/upload",
        headers=auth_headers,
        files={"file": ("resume.pdf", b"not a real pdf", "application/pdf")},
    )
    assert response.status_code == 400
    assert "extension" in response.json()["detail"]


def test_upload_rejects_corrupt_pdf(client: TestClient, auth_headers: dict) -> None:
    payload = b"%PDF-1.7\n%garbage that pymupdf cannot open"
    response = client.post(
        "/api/resumes/upload",
        headers=auth_headers,
        files={"file": ("resume.pdf", payload, "application/pdf")},
    )
    assert response.status_code == 400
    assert "parse" in response.json()["detail"].lower()


def test_upload_rejects_empty_file(client: TestClient, auth_headers: dict) -> None:
    response = client.post(
        "/api/resumes/upload",
        headers=auth_headers,
        files={"file": ("resume.pdf", b"", "application/pdf")},
    )
    assert response.status_code == 400


def test_oversized_body_is_rejected(client: TestClient, auth_headers: dict) -> None:
    response = client.get(
        "/api/resumes",
        headers={**auth_headers, "Content-Length": str(50 * 1024 * 1024)},
    )
    assert response.status_code == 413
    assert response.json()["code"] == "payload_too_large"


def test_upload_valid_pdf_is_listed_fetched_and_deleted(
    client: TestClient, auth_headers: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(RESUME_SERVICE_LLM, fake_resume_extraction)
    payload = make_pdf()

    created = client.post(
        "/api/resumes/upload",
        headers=auth_headers,
        files={"file": ("resume.pdf", payload, "application/pdf")},
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["filename"] == "resume.pdf"
    assert body["structured_data"]["skills"] == ["Python", "FastAPI"]
    assert body["analysis_status"] == "ready"
    assert body["ocr_used"] is False
    resume_id = body["id"]

    listed = client.get("/api/resumes", headers=auth_headers)
    assert listed.status_code == 200
    summaries = listed.json()
    summary = next(item for item in summaries if item["id"] == resume_id)
    assert summary["analysis_status"] == "ready"
    assert summary["has_preview"] is True
    assert all("raw_text" not in item for item in summaries)

    detail = client.get(f"/api/resumes/{resume_id}", headers=auth_headers)
    assert detail.status_code == 200
    assert detail.json()["raw_text"]
    assert detail.json()["analysis"]["status"] == "ready"
    # Storage identifiers are implementation details and never leave the API.
    assert "storage_key" not in detail.json()
    assert "original_file_id" not in detail.json()
    assert all("file_id" not in preview for preview in detail.json()["previews"])

    text = client.get(f"/api/resumes/{resume_id}/text", headers=auth_headers)
    assert text.status_code == 200
    assert "Jane Doe" in text.json()["text"]

    download = client.get(f"/api/resumes/{resume_id}/file", headers=auth_headers)
    assert download.status_code == 200
    assert download.content == payload
    assert "resume.pdf" in download.headers["content-disposition"]
    assert download.headers["x-content-type-options"] == "nosniff"

    preview = client.get(f"/api/resumes/{resume_id}/preview/1", headers=auth_headers)
    assert preview.status_code == 200
    assert preview.headers["content-type"] == "image/png"
    assert preview.content.startswith(b"\x89PNG")

    missing_preview = client.get(
        f"/api/resumes/{resume_id}/preview/99", headers=auth_headers
    )
    assert missing_preview.status_code == 404

    assert client.delete(f"/api/resumes/{resume_id}", headers=auth_headers).status_code == 204
    assert client.get(f"/api/resumes/{resume_id}", headers=auth_headers).status_code == 404
    assert client.get(f"/api/resumes/{resume_id}/file", headers=auth_headers).status_code == 404


def test_docx_upload_preserves_paragraphs_tables_and_text(
    client: TestClient, auth_headers: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(RESUME_SERVICE_LLM, fake_resume_extraction)

    created = client.post(
        "/api/resumes/upload",
        headers=auth_headers,
        files={
            "file": (
                "resume.docx",
                make_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["analysis_status"] == "ready"
    assert body["filename"] == "resume.docx"

    text = client.get(f"/api/resumes/{body['id']}/text", headers=auth_headers)
    assert text.status_code == 200
    extracted = text.json()["text"]
    assert "Jane Doe" in extracted
    assert "Company | Acme" in extracted
    assert "Python, FastAPI" in extracted


def test_duplicate_upload_reuses_the_stored_resume(
    client: TestClient, auth_headers: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(RESUME_SERVICE_LLM, fake_resume_extraction)
    payload = make_pdf("Jane Doe - duplicate upload check.")

    first = client.post(
        "/api/resumes/upload",
        headers=auth_headers,
        files={"file": ("resume.pdf", payload, "application/pdf")},
    )
    assert first.status_code == 201, first.text

    second = client.post(
        "/api/resumes/upload",
        headers=auth_headers,
        files={"file": ("resume.pdf", payload, "application/pdf")},
    )
    assert second.status_code == 200, second.text
    assert second.json()["deduplicated"] is True
    assert second.json()["id"] == first.json()["id"]


def test_large_resumes_are_extracted_in_bounded_chunks(
    client: TestClient, auth_headers: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Long documents are split so no single provider request exceeds the budget."""
    prompts: list[str] = []
    calls: list[dict] = []

    def huge_pages(file_bytes: bytes):
        return [{"page": 0, "text": "Experience: backend services.\n" * 5000}]

    async def capturing_extraction(prompt: str, schema, **kwargs):
        prompts.append(prompt)
        calls.append(kwargs)
        return schema(skills=[], experience=[], projects=[], education=[])

    monkeypatch.setattr(RESUME_PARSE_PAGES, huge_pages)
    monkeypatch.setattr(RESUME_SERVICE_LLM, capturing_extraction)

    created = client.post(
        "/api/resumes/upload",
        headers=auth_headers,
        files={"file": ("resume.pdf", make_pdf(), "application/pdf")},
    )
    assert created.status_code == 201, created.text
    assert len(prompts) > 1
    assert max(len(prompt) for prompt in prompts) < (
        settings.resume_extraction_chunk_chars + 15_000
    )
    assert all(call.get("system_prompt") for call in calls)


def test_upload_is_queued_when_inline_processing_is_disabled(
    client: TestClient,
    auth_headers: dict,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "resume_processing_inline", False)
    monkeypatch.setattr(RESUME_SERVICE_LLM, fake_resume_extraction)

    created = client.post(
        "/api/resumes/upload",
        headers=auth_headers,
        files={"file": ("resume.pdf", make_pdf(), "application/pdf")},
    )
    assert created.status_code == 202, created.text
    body = created.json()
    assert body["analysis_status"] == "pending"
    resume_id = body["id"]

    async def run_handler(db) -> None:
        queued = await db.background_jobs.find_one(
            {"type": RESUME_PROCESS_JOB, "payload.resume_id": resume_id}
        )
        assert queued is not None
        handler = job_worker.handlers()[RESUME_PROCESS_JOB]
        await handler(db, queued["payload"])

    _with_mongo(run_handler)

    detail = client.get(f"/api/resumes/{resume_id}", headers=auth_headers)
    assert detail.status_code == 200
    assert detail.json()["analysis"]["status"] == "ready"


def test_upload_survives_llm_failure_and_can_be_retried(
    client: TestClient, auth_headers: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def boom(prompt: str, schema, **kwargs):
        raise RuntimeError("llm is down")

    monkeypatch.setattr(RESUME_SERVICE_LLM, boom)

    response = client.post(
        "/api/resumes/upload",
        headers=auth_headers,
        files={"file": ("resume.pdf", make_pdf(), "application/pdf")},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["analysis_status"] == "failed"
    resume_id = body["id"]

    # The original, previews and full text remain available for manual use.
    detail = client.get(f"/api/resumes/{resume_id}", headers=auth_headers)
    assert detail.status_code == 200
    assert detail.json()["raw_text"]
    assert detail.json()["previews"]
    assert detail.json()["analysis"]["error_code"] == "llm_failed"

    monkeypatch.setattr(RESUME_SERVICE_LLM, fake_resume_extraction)
    retried = client.post(f"/api/resumes/{resume_id}/analyze", headers=auth_headers)
    assert retried.status_code == 200, retried.text
    assert retried.json()["analysis_status"] == "ready"
    assert (
        client.get(f"/api/resumes/{resume_id}", headers=auth_headers)
        .json()["structured_data"]["skills"]
        == ["Python", "FastAPI"]
    )


def test_get_resume_with_invalid_id_is_404(client: TestClient, auth_headers: dict) -> None:
    response = client.get("/api/resumes/not-an-object-id", headers=auth_headers)
    assert response.status_code == 404
    assert response.json() == {"detail": "Resume not found", "code": "not_found"}


def test_delete_resume_with_invalid_id_is_404(client: TestClient, auth_headers: dict) -> None:
    response = client.delete("/api/resumes/not-an-object-id", headers=auth_headers)
    assert response.status_code == 404


def test_delete_resume_of_another_user_is_404(
    client: TestClient, auth_headers: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(RESUME_SERVICE_LLM, fake_resume_extraction)
    created = client.post(
        "/api/resumes/upload",
        headers=auth_headers,
        files={"file": ("resume.pdf", make_pdf(), "application/pdf")},
    )
    resume_id = created.json()["id"]

    # A second account must not be able to read or delete someone else's resume.
    other_email = "other.user@example.com"
    register = client.post(
        "/api/auth/register",
        json={"name": "Other User", "email": other_email, "password": "strongpassword1"},
    )
    login = client.post(
        "/api/auth/login",
        json={"email": other_email, "password": "strongpassword1"},
    )
    assert register.status_code == 201
    other = {"Authorization": f"Bearer {login.json()['access_token']}"}

    assert client.get(f"/api/resumes/{resume_id}", headers=other).status_code == 404
    assert client.delete(f"/api/resumes/{resume_id}", headers=other).status_code == 404
    # The owner's copy is untouched.
    assert client.get(f"/api/resumes/{resume_id}", headers=auth_headers).status_code == 200
    assert client.delete(f"/api/resumes/{resume_id}", headers=auth_headers).status_code == 204


def test_list_resumes_returns_summaries(client: TestClient, auth_headers: dict) -> None:
    response = client.get("/api/resumes", headers=auth_headers)
    assert response.status_code == 200
    assert isinstance(response.json(), list)


def test_job_description_requires_company(client: TestClient, auth_headers: dict) -> None:
    response = client.post(
        "/api/job-descriptions",
        headers=auth_headers,
        json={"title": "Backend Engineer", "raw_text": "x" * 60},
    )
    assert response.status_code == 422


def test_job_description_requires_auth(client: TestClient) -> None:
    assert client.get("/api/job-descriptions").status_code == 401
    assert client.post("/api/job-descriptions", json={}).status_code == 401


def test_job_description_with_invalid_id_is_404(client: TestClient, auth_headers: dict) -> None:
    response = client.get("/api/job-descriptions/nope", headers=auth_headers)
    assert response.status_code == 404


def test_job_description_create_list_get_delete(
    client: TestClient, auth_headers: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def fake_extract(prompt: str, schema, **kwargs):
        return schema(
            role_title="Backend Engineer",
            required_skills=["FastAPI"],
            responsibilities=["Build APIs"],
        )

    monkeypatch.setattr(JD_SERVICE_LLM, fake_extract)

    payload = {
        "title": "Backend Engineer",
        "company": "Acme",
        "raw_text": "Own our FastAPI services and MongoDB collections end to end.",
    }
    created = client.post("/api/job-descriptions", headers=auth_headers, json=payload)
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["title"] == "Backend Engineer"
    assert body["structured_data"]["required_skills"] == ["FastAPI"]
    jd_id = body["id"]

    listed = client.get("/api/job-descriptions", headers=auth_headers)
    assert listed.status_code == 200
    summaries = listed.json()
    assert any(item["id"] == jd_id for item in summaries)
    assert all("raw_text" not in item for item in summaries)
    assert all("snippet" in item for item in summaries)

    detail = client.get(f"/api/job-descriptions/{jd_id}", headers=auth_headers)
    assert detail.status_code == 200
    assert detail.json()["raw_text"] == payload["raw_text"]

    assert client.delete(f"/api/job-descriptions/{jd_id}", headers=auth_headers).status_code == 204
    assert client.get(f"/api/job-descriptions/{jd_id}", headers=auth_headers).status_code == 404


def test_job_description_saved_when_llm_unavailable(
    client: TestClient, auth_headers: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def boom(prompt: str, schema, **kwargs):
        raise RuntimeError("llm is down")

    monkeypatch.setattr(JD_SERVICE_LLM, boom)

    response = client.post(
        "/api/job-descriptions",
        headers=auth_headers,
        json={"title": "SRE", "company": "Acme", "raw_text": "Keep the cluster healthy."},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["structured_data"]["role_title"] == "SRE"
    assert body["structured_data"]["required_skills"] == []


def test_job_description_list_handles_legacy_documents(
    client: TestClient, auth_headers: dict
) -> None:
    """Older documents have no `snippet` and may store null lists."""
    me = client.get("/api/auth/me", headers=auth_headers).json()
    insert_job_description(
        {
            "user_id": me["id"],
            "title": "Legacy Role",
            "company": "Old Co",
            "raw_text": "A posting saved before the snippet field existed.",
            "structured_data": {"responsibilities": None},
            "created_at": datetime.now(UTC),
        }
    )

    response = client.get("/api/job-descriptions", headers=auth_headers)
    assert response.status_code == 200
    legacy = next(item for item in response.json() if item["title"] == "Legacy Role")
    assert legacy["snippet"] == ""


def test_chunk_merge_preserves_link_objects() -> None:
    from app.services.resume_processing import _merge_objects

    merged = _merge_objects(
        [
            {
                "name": "Portfolio",
                "links": [{"label": "GitHub", "url": "https://github.com/a"}],
            },
            {
                "name": "Portfolio",
                "links": [{"label": "Demo", "url": "https://example.com/demo"}],
            },
        ],
        ("name",),
    )
    assert len(merged) == 1
    assert merged[0]["links"] == [
        {"label": "GitHub", "url": "https://github.com/a"},
        {"label": "Demo", "url": "https://example.com/demo"},
    ]


def test_missing_resume_sections_default_to_empty_lists() -> None:
    """A section the model omits must not fail extraction (e.g. no education)."""
    from app.schemas.resume import ResumeStructuredData

    parsed = ResumeStructuredData.model_validate(
        {"skills": ["Python"], "experience": [], "projects": []}
    )
    assert parsed.education == []
    assert parsed.skills == ["Python"]

    # The empty object is valid too -- every section is optional.
    assert ResumeStructuredData.model_validate({}).education == []
