import io

from docx import Document

from app.services.parsers.text_cleaner import clean_text


def parse_docx(file_bytes: bytes) -> str:
    """Extract text (paragraphs *and* tables) from a DOCX file."""
    doc = Document(io.BytesIO(file_bytes))
    blocks: list[str] = [para.text for para in doc.paragraphs if para.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))

    return clean_text("\n".join(blocks))
